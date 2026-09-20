"""
Batch-based exam report API endpoints for PROCTIFY.
"""
import mysql.connector
from flask import Blueprint, jsonify, request, session, send_file
from datetime import datetime
import os
import io

batch_reports_bp = Blueprint("batch_reports", __name__)


def _get_db():
    return mysql.connector.connect(
        host=os.environ.get("PROCTIFY_DB_HOST", "localhost"),
        user=os.environ.get("PROCTIFY_DB_USER", "root"),
        password=os.environ.get("PROCTIFY_DB_PASSWORD", "aasmaan@14"),
        database=os.environ.get("PROCTIFY_DB_NAME", "proctify_db"),
    )


def _teacher_logged_in():
    return bool(session.get("teacher_id"))


@batch_reports_bp.route('/api/batch-reports', methods=['GET'])
def api_batch_reports():
    if not _teacher_logged_in():
        return jsonify({"success": False, "error": "Teacher login required."}), 401

    connection = None
    cursor = None
    try:
        connection = _get_db()
        cursor = connection.cursor(dictionary=True)
        cursor.execute('''
            SELECT
                ea.batch_id,
                b.batch_name,
                ea.exam_id,
                e.exam_name,
                COUNT(*) AS student_count
            FROM exam_assignments ea
            INNER JOIN batches b ON b.batch_id = ea.batch_id
            INNER JOIN exams e ON e.exam_id = ea.exam_id
            WHERE ea.batch_id IS NOT NULL
            GROUP BY ea.batch_id, ea.exam_id
            ORDER BY b.batch_name, e.exam_name''')
        rows = cursor.fetchall()
        reports = []
        for row in rows:
            reports.append({
                "batch_id": row["batch_id"],
                "batch_name": row["batch_name"],
                "exam_id": row["exam_id"],
                "exam_name": row["exam_name"],
                "student_count": int(row.get("student_count") or 0),
            })
        return jsonify({"success": True, "reports": reports})
    except Exception as error:
        print("Batch reports error:", error)
        return jsonify({"success": False, "error": str(error), "reports": []}), 500
    finally:
        if cursor is not None:
            cursor.close()
        if connection is not None:
            connection.close()


def _fetch_batch_students(cursor, batch_id, exam_id):
    cursor.execute("SELECT exam_name FROM exams WHERE exam_id = %s LIMIT 1", (str(exam_id),))
    exam_row = cursor.fetchone()
    exam_name_val = exam_row["exam_name"] if exam_row else ""

    cursor.execute('''
            SELECT
                s.student_id,
                s.student_name,
                es.session_id,
                es.status AS session_status,
                es.final_trust_score,
                es.final_risk_level,
                sub.submission_id,
                sub.score,
                sub.teacher_marks,
                sub.trust_score AS submission_trust,
                sub.total_questions,
                sub.status AS submission_status,
                sub.evaluation_type,
                sub.teacher_feedback,
                sub.submitted_at,
                (
                    SELECT COUNT(*)
                    FROM violations v
                    WHERE v.session_id = es.session_id
                ) AS violation_count,
                (
                    SELECT COUNT(*)
                    FROM evidence ev
                    WHERE ev.session_id = es.session_id
                ) AS evidence_count
            FROM batch_members bm
            INNER JOIN students s
                ON s.student_id = bm.student_id
            LEFT JOIN exam_assignments ea
                ON ea.student_id = s.student_id
                AND ea.exam_id = %s
                AND ea.batch_id = %s
            LEFT JOIN exam_sessions es
                ON es.student_id = s.student_id
                AND es.exam_name = %s
            LEFT JOIN exam_submissions sub
                ON sub.student_id = s.student_id
                AND sub.exam_id = %s
            WHERE bm.batch_id = %s
            ORDER BY s.student_id''', (
        str(exam_id), batch_id, exam_name_val, str(exam_id), batch_id
    ))
    students = cursor.fetchall()

    for stu in students:
        for key, val in list(stu.items()):
            if isinstance(val, datetime):
                stu[key] = val.isoformat()

        trust = stu.get("final_trust_score")
        if trust is None:
            trust = stu.get("submission_trust")
        stu["trust_score_final"] = trust

        session_status = (stu.get("session_status") or "").upper()
        submission_status = (stu.get("submission_status") or "").upper()
        score = stu.get("score")
        teacher_marks = stu.get("teacher_marks")

        if session_status == "TERMINATED":
            stu["display_status"] = "AUTO TERMINATED"
        elif submission_status == "EVALUATED":
            effective_marks = teacher_marks if teacher_marks is not None else score
            if effective_marks is not None:
                stu["display_status"] = "PASS" if float(effective_marks) >= 50 else "FAIL"
            else:
                stu["display_status"] = "EVALUATED"
        elif submission_status in ("SUBMITTED", "PENDING_REVIEW"):
            stu["display_status"] = "PENDING REVIEW"
        elif submission_status == "AUTO_GRADING":
            stu["display_status"] = "AUTO GRADING"
        elif session_status in ("ONGOING", "COMPLETED"):
            if stu.get("submission_id") is None:
                stu["display_status"] = "EXAM NOT ATTEMPTED"
            else:
                stu["display_status"] = "INCOMPLETE"
        elif stu.get("submission_id") is None:
            stu["display_status"] = "EXAM NOT ATTEMPTED"
        else:
            stu["display_status"] = submission_status or session_status or "EXAM NOT ATTEMPTED"

        effective_marks = teacher_marks if teacher_marks is not None else score
        if stu["display_status"] == "EXAM NOT ATTEMPTED":
            stu["marks_display"] = None
            stu["pass_fail"] = None
        elif effective_marks is not None:
            stu["marks_display"] = effective_marks
            stu["pass_fail"] = "PASS" if float(effective_marks) >= 50 else "FAIL"
        else:
            stu["marks_display"] = None
            stu["pass_fail"] = None

        stu.pop("submission_trust", None)

    return exam_name_val, students


@batch_reports_bp.route('/api/batch-reports/<int:batch_id>/<exam_id>', methods=['GET'])
def api_batch_report_detail(batch_id, exam_id):
    if not _teacher_logged_in():
        return jsonify({"success": False, "error": "Teacher login required."}), 401

    connection = None
    cursor = None
    try:
        connection = _get_db()
        cursor = connection.cursor(dictionary=True)

        cursor.execute("SELECT batch_name FROM batches WHERE batch_id = %s LIMIT 1", (batch_id,))
        batch_row = cursor.fetchone()
        if not batch_row:
            return jsonify({"success": False, "error": "Batch not found."}), 404

        exam_name_val, students = _fetch_batch_students(cursor, batch_id, exam_id)

        return jsonify({
            "success": True,
            "batch_id": batch_id,
            "batch_name": batch_row["batch_name"],
            "exam_id": str(exam_id),
            "exam_name": exam_name_val,
            "students": students,
            "total_students": len(students),
        })
    except Exception as error:
        print("Batch report detail error:", error)
        return jsonify({"success": False, "error": str(error)}), 500
    finally:
        if cursor is not None:
            cursor.close()
        if connection is not None:
            connection.close()


@batch_reports_bp.route('/api/batch-reports/export/excel', methods=['POST'])
def export_batch_report_excel():
    if not _teacher_logged_in():
        return jsonify({"success": False, "error": "Teacher login required."}), 401

    data = request.get_json(silent=True) or {}
    batch_id = data.get("batch_id")
    exam_id = data.get("exam_id")

    if batch_id is None or exam_id is None:
        return jsonify({"success": False, "error": "batch_id and exam_id required."}), 400

    connection = None
    cursor = None
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment

        connection = _get_db()
        cursor = connection.cursor(dictionary=True)

        cursor.execute("SELECT batch_name FROM batches WHERE batch_id = %s LIMIT 1", (int(batch_id),))
        batch_row = cursor.fetchone()
        if not batch_row:
            return jsonify({"success": False, "error": "Batch not found."}), 404

        batch_name = batch_row["batch_name"]
        exam_name_val, students = _fetch_batch_students(cursor, int(batch_id), exam_id)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Batch Report"

        ws.cell(row=1, column=1, value="Batch:").font = Font(bold=True)
        ws.cell(row=1, column=2, value=batch_name)
        ws.cell(row=2, column=1, value="Exam:").font = Font(bold=True)
        ws.cell(row=2, column=2, value=exam_name_val)
        ws.cell(row=3, column=1, value="Date:").font = Font(bold=True)
        ws.cell(row=3, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M"))

        headers = [
            "Student Name", "Student ID", "Exam Name",
            "Marks", "Total Marks", "Status",
            "Trust Score", "Risk Level",
            "Violations", "Evidence Count",
            "Evaluation", "Termination Reason",
        ]
        header_font = Font(bold=True)
        header_row = 5
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=h)
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        row_num = header_row + 1
        for stu in students:
            session_status = (stu.get("session_status") or "").upper()
            submission_status = (stu.get("submission_status") or "").upper()
            score = stu.get("score")
            teacher_marks = stu.get("teacher_marks")
            trust = stu.get("final_trust_score")
            if trust is None:
                trust = stu.get("submission_trust")
            effective_marks = teacher_marks if teacher_marks is not None else score

            if session_status == "TERMINATED":
                display_status = "AUTO TERMINATED"
                termination_reason = "Trust Score reached 0"
            elif submission_status == "EVALUATED":
                if effective_marks is not None:
                    display_status = "PASS" if float(effective_marks) >= 50 else "FAIL"
                else:
                    display_status = "EVALUATED"
                termination_reason = ""
            elif submission_status in ("SUBMITTED", "PENDING_REVIEW"):
                display_status = "PENDING REVIEW"
                termination_reason = ""
            elif submission_status == "AUTO_GRADING":
                display_status = "AUTO GRADING"
                termination_reason = ""
            elif session_status in ("ONGOING", "COMPLETED"):
                if stu.get("submission_id") is None:
                    display_status = "EXAM NOT ATTEMPTED"
                else:
                    display_status = "INCOMPLETE"
                termination_reason = ""
            elif stu.get("submission_id") is None:
                display_status = "EXAM NOT ATTEMPTED"
                termination_reason = ""
            else:
                display_status = submission_status or session_status or "EXAM NOT ATTEMPTED"
                termination_reason = ""

            values = [
                stu.get("student_name", ""),
                stu.get("student_id", ""),
                exam_name_val,
                effective_marks if effective_marks is not None else "-",
                100,
                display_status,
                trust if trust is not None else "-",
                stu.get("final_risk_level", "-"),
                int(stu.get("violation_count") or 0),
                int(stu.get("evidence_count") or 0),
                stu.get("evaluation_type", "-"),
                termination_reason,
            ]
            for col, val in enumerate(values, 1):
                ws.cell(row=row_num, column=col, value=val)
            row_num += 1

        for col_cells in ws.columns:
            max_len = 0
            col_letter = col_cells[0].column_letter
            for cell in col_cells:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 3, 30)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        safe_name = batch_name.replace(" ", "_").replace("/", "_")
        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"{safe_name} Report.xlsx",
        )
    except Exception as error:
        print("Batch report Excel export error:", error)
        return jsonify({"success": False, "error": str(error)}), 500
    finally:
        if cursor is not None:
            cursor.close()
        if connection is not None:
            connection.close()
